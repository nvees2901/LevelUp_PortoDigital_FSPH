"""
docx_generator.py — Geração de DOCX do Termo de Referência (python-docx).

Espelha a estrutura do PDF (cabeçalho institucional, título, seções a partir
do Markdown, assinaturas) em um documento Word editável. Reutiliza
`clean_tr_content` para remover cercas/preâmbulo/fecho conversacional.
"""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.services.pdf_generator import clean_tr_content
from app.utils.logging import get_logger

logger = get_logger(__name__)

# Termo padronizado em letras pretas.
BRAND = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x00, 0x00, 0x00)
FONT = "Calibri"


def _runs_markdown(paragraph, text: str) -> None:
    """Adiciona texto com **negrito**/*itálico* como runs ao parágrafo."""
    # remove marcação que não viramos run (código/links) -> texto puro
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", text)
    # tokeniza **bold** e *italic*
    parts = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|__.+?__|_[^_]+?_)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            run.text = part[2:-2]
            run.bold = True
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            run.text = part[1:-1]
            run.italic = True
        else:
            run.text = part
        run.font.name = FONT


class DocxGeneratorService:

    @classmethod
    def generate_term_docx(cls, term_data: dict) -> bytes:
        doc = Document()

        normal = doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(11)

        # --- Cabeçalho institucional ---
        cls._centered(doc, "GOVERNO DO ESTADO DE SERGIPE", size=10, color=GRAY)
        cls._centered(doc, "FUNDAÇÃO DE SAÚDE PARREIRAS HORTA — FSPH", size=13, bold=True, color=BRAND)
        cls._centered(doc, "Comissão de Licitação | Lei nº 14.133/2021", size=9, color=GRAY)
        doc.add_paragraph()

        # --- Título ---
        cls._centered(doc, "TERMO DE REFERÊNCIA", size=16, bold=True, color=BRAND)
        objeto = (term_data.get("title") or "").strip()
        if objeto:
            cls._centered(doc, objeto, size=11, color=GRAY)
        doc.add_paragraph()

        # --- Corpo (markdown) ---
        content = clean_tr_content(term_data.get("content") or "")
        if content.strip():
            cls._render_markdown(doc, content)
        else:
            doc.add_paragraph("Conteúdo do termo não disponível.")

        # --- Local e data + assinaturas ---
        doc.add_paragraph()
        ld = doc.add_paragraph(term_data.get("local_data") or "")
        ld.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()
        doc.add_paragraph()
        cls._signature(doc, term_data)

        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()
        buf.close()
        logger.info("DOCX gerado: term_id=%s size=%d bytes", term_data.get("id"), len(data))
        return data

    # ------------------------------------------------------------------ #
    @staticmethod
    def _centered(doc, text, size=11, bold=False, color=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.name = FONT
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return p

    @classmethod
    def _render_markdown(cls, doc, md: str) -> None:
        lines = md.split("\n")
        i = 0
        para_buf: list[str] = []

        def flush():
            if para_buf:
                text = " ".join(s.strip() for s in para_buf).strip()
                if text:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    _runs_markdown(p, text)
                para_buf.clear()

        while i < len(lines):
            line = lines[i].rstrip()
            s = line.strip()
            if not s:
                flush(); i += 1; continue

            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                flush()
                level = len(m.group(1))
                htext = m.group(2).strip().strip("*").strip()
                if level == 1 and re.search(r"termo de refer", htext, re.I) and len(htext) < 40:
                    i += 1; continue
                h = doc.add_heading(level=min(level, 3))
                run = h.add_run(htext)
                run.font.name = FONT
                run.font.color.rgb = BRAND
                i += 1; continue

            if "|" in s:
                block = []
                while i < len(lines) and "|" in lines[i] and lines[i].strip():
                    block.append(lines[i]); i += 1
                if cls._add_table(doc, block):
                    continue
                para_buf.extend(block); continue

            mb = re.match(r"^[-*+]\s+(.*)$", s)
            if mb:
                flush()
                p = doc.add_paragraph(style="List Bullet")
                _runs_markdown(p, mb.group(1))
                i += 1; continue

            mn = re.match(r"^(\d{1,2}[.)]|[a-zA-Z][.)])\s+(.*)$", s)
            if mn:
                flush()
                p = doc.add_paragraph()
                _runs_markdown(p, f"{mn.group(1)} {mn.group(2)}")
                i += 1; continue

            para_buf.append(s); i += 1
        flush()

    @staticmethod
    def _add_table(doc, block: list[str]) -> bool:
        def split_row(x: str):
            x = x.strip().strip("|")
            return [c.strip() for c in x.split("|")]
        rows = [split_row(b) for b in block if b.strip()]
        rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
        if not rows:
            return False
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        table = doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = table.cell(ri, ci)
                cell.text = ""
                _runs_markdown(cell.paragraphs[0], val)
                if ri == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        return True

    @staticmethod
    def _signature(doc, term_data: dict) -> None:
        nome = (term_data.get("elaborador_nome") or "").strip() or "Responsável pela elaboração"
        matricula = (term_data.get("elaborador_matricula") or "").strip()
        setor = (term_data.get("elaborador_setor") or "").strip()
        sub = f"Matrícula: {matricula}" + (f" — {setor}" if setor else "") if matricula else "Cargo / Matrícula"

        aut_nome = (term_data.get("autoridade_nome") or "").strip()
        aut_cargo = (term_data.get("autoridade_cargo") or "").strip()
        aut_top = aut_nome or "Autoridade competente"
        aut_role = "Autoridade competente" if aut_nome else "Cargo / Função"
        aut_sub = aut_cargo or ""

        t = doc.add_table(rows=4, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        data = [
            ("_" * 34, "_" * 34),
            (nome, aut_top),
            ("Responsável pela elaboração", aut_role),
            (sub, aut_sub),
        ]
        for ri, (l, r) in enumerate(data):
            for ci, val in enumerate((l, r)):
                cell = t.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.font.name = FONT
                run.font.size = Pt(10)
                if ri == 1:
                    run.bold = True
